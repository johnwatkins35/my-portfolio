"""Generate SysDesk write-up practice template. Run: python build_writeup_template.py"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "SysDesk-Writeup-Practice-Template.docx"

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
GOLD = RGBColor(0x8A, 0x6D, 0x12)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x4A, 0x55, 0x63)

TICKETS = [
    ("INC-2401", "Help desk / Northwind", "Jordan Lee lockout", "Confirm identity. Unlock jlee. Do not escalate."),
    ("INC-2402", "Help desk / Northwind", "Priya password / VPN", "Reset password. Read it on chat. Do not email it."),
    ("INC-2403", "Help desk / Northwind", "Chris warehouse share", "Fix DNS to 10.20.0.11. Add VPN-Users."),
    ("INC-2404", "Help desk / Northwind", "Hannah spooler", "RDP NW-WS-0881. Start Print Spooler."),
    ("INC-2405", "Help desk / Northwind", "Elena contractor rehire", "Escalate to HR. Do not enable the account."),
    ("INC-2406", "Help desk / Northwind", "Kenji VPN hang", "RDP. Clear stale VPN profile."),
    ("INC-2398", "Help desk / Northwind", "Local / Domain Admin ask", "Escalate. Do not grant admin."),
    ("INC-2501", "MSP / Harbor Dental", "Maya clinic no websites", "Claim. Vault hd. RDP HD-WS-14. ipconfig /renew."),
    ("INC-2502", "MSP / Riverton", "Maya Chen teacher lockout", "Unlock mchenrs. Not Harbor's Maya (mchen)."),
    ("INC-2503", "MSP / Maple CU", "Asha Defender off", "Claim. Vault mc. RDP. Start Defender."),
    ("INC-2504", "MSP / Oak & Pine", "Blair wants Domain Admin", "Claim. Escalate. Do not grant DA."),
    ("INC-2505", "MSP / Peak", "Santiago Ortiz lockout", "Unlock sortiz. Not Sam Ortiz at Oak & Pine."),
    ("CLD-1001", "Cloud", "web-prod stopped", "Start the VM."),
    ("CLD-1002", "Cloud", "Leaked deploy-bot key", "Disable AKIA-OLD1. Do not paste a new secret."),
    ("CLD-1003", "Cloud", "SSH open on sg-web", "Lock port 22 to 10.20.0.0/16."),
    ("ACS-2001", "Access", "Chris dock deny", "Grant B-4412 door WH-DOCK."),
    ("ACS-2002", "Access", "Exec controller down", "Reboot CTRL-07."),
    ("ACS-2003", "Access", "Elena badge after term", "Disable B-0771. AD disable is not enough."),
    ("SOC-3001", "SOC", "Ransom on Kenji PC", "Isolate NW-WS-3304."),
    ("SOC-3002", "SOC", "jlee failed-logon burst", "False positive. Do not isolate the AP clerk."),
    ("SOC-3003", "SOC", "Riley macro zip", "Isolate NW-LT-0007."),
    ("NOC-4001", "NOC", "Warehouse uplink down", "Bounce CIR-WH."),
    ("NOC-4002", "NOC", "ExpressRoute degraded", "Escalate. Do not bounce CIR-AZ."),
    ("DBA-5001", "DBA", "Orders blocked", "Kill SPID 88. Do not restart SQL."),
    ("DBA-5002", "DBA", "Nightly ETL failed", "Rerun JOB-NIGHT."),
    ("DBA-5003", "DBA", "Add an index today", "Escalate. That is a change."),
    ("CHG-6001", "Change", "Standard SG /32", "Approve. Has backout and window."),
    ("CHG-6002", "Change", "DC patch, no backout", "Reject."),
]


def set_run(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_p(doc, text, size=11, bold=False, color=INK, space_after=8, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
    return h


def field_row(table, label, hint="", tall=False):
    row = table.add_row().cells
    row[0].text = ""
    row[1].text = ""
    p0 = row[0].paragraphs[0]
    r0 = p0.add_run(label)
    set_run(r0, size=10, bold=True, color=NAVY)
    if hint:
        p_hint = row[0].add_paragraph()
        rh = p_hint.add_run(hint)
        set_run(rh, size=8, color=MUTED, italic=True)
    p1 = row[1].paragraphs[0]
    p1.add_run("")
    blanks = 4 if tall else 2
    for _ in range(blanks):
        row[1].add_paragraph()
    return row


def blank_form(doc, title="Ticket write-up (copy this page)"):
    heading(doc, title, 1)
    add_p(
        doc,
        "Fill this after you close the ticket in SysDesk. Write what you actually did. Leave a line blank rather than inventing.",
        size=10,
        color=MUTED,
        italic=True,
        space_after=10,
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.columns[0].width = Inches(2.1)
    table.columns[1].width = Inches(4.6)
    hdr = table.rows[0].cells
    hdr[0].text = ""
    hdr[1].text = ""
    hr = hdr[0].paragraphs[0].add_run("Field")
    set_run(hr, size=10, bold=True, color=NAVY)
    hr2 = hdr[1].paragraphs[0].add_run("Your notes")
    set_run(hr2, size=10, bold=True, color=NAVY)

    rows = [
        ("Date / session", "e.g. 30 Aug 2026 · practice shift 2", False),
        ("Ticket ID + desk + client", "INC-     · Help desk / Cloud / Access / …", False),
        ("Channel / priority", "Voice · email · chat     P1 / P2 / P3", False),
        ("Caller + SAM + host", "Name · sam · workstation · tenant", False),
        ("Situation (2–3 sentences)", "What they said vs what was actually broken", True),
        ("Identity / asset check", "How you confirmed the person and the device", True),
        ("Actions taken", "Tools. Exact change (unlock, vault, isolate…)", True),
        ("Escalate or resolve?", "Why it was in scope — or why you refused", True),
        ("What you told the user", "No passwords in email. Timebox. Next step.", True),
        ("Work notes (type 40+ chars)", "Write here. Or paste from SysDesk, then clean up.", True),
        ("Outcome", "Verified / Unverified / Escalated", False),
        ("Skills this demonstrates", "Pick 2–3 from the Skills page", False),
        ("STAR — Situation", "Set the scene in two sentences", True),
        ("STAR — Task", "What you owned", True),
        ("STAR — Action", "What you did, in order", True),
        ("STAR — Result", "Verified close / user working / handed to T2", True),
        ("Interview one-liner", "One sentence you can say out loud", True),
        ("What I would do differently", "Optional. Honest. Short.", False),
    ]
    for label, hint, tall in rows:
        field_row(table, label, hint, tall=tall)

    add_p(doc, "", space_after=4)
    add_p(
        doc,
        "This is a simulated SysDesk / Northwind MSP lab. Not a job at Northwind, AWS, or a real clinic/bank.",
        size=9,
        color=MUTED,
        italic=True,
    )


def example_form(doc):
    heading(doc, "Worked example — INC-2401 (do not copy as if it were your job)", 1)
    add_p(
        doc,
        "Use this as the bar for length and honesty. Write your own tickets in the same shape.",
        size=10,
        color=MUTED,
        italic=True,
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.columns[0].width = Inches(2.1)
    table.columns[1].width = Inches(4.6)
    hdr = table.rows[0].cells
    hdr[0].paragraphs[0].add_run("Field").bold = True
    hdr[1].paragraphs[0].add_run("Example").bold = True
    example = [
        ("Ticket ID", "INC-2401 · Help desk / Northwind · voice · P2"),
        ("Caller + SAM + host", "Jordan Lee · jlee · NW-WS-1042"),
        (
            "Situation",
            "AP clerk locked out after a Friday password change. Typed the old password. Close is this morning.",
        ),
        (
            "Identity / asset check",
            "Matched caller name, SAM jlee, and NW-WS-1042 in Directory. Account showed Locked, enabled.",
        ),
        (
            "Actions taken",
            "Confirmed identity on the ticket. Unlocked jlee in the simulated AD. Stayed on the voice channel. Documented lockout source as bad password, not an attack.",
        ),
        (
            "Escalate or resolve?",
            "Resolved. Simple lockout is Tier-1. I would escalate if the lockout source looked like a spray or the account was disabled.",
        ),
        (
            "What you told the user",
            "Owned the ticket, gave a short timebox, unlocked, asked them to sign in with the Friday password. Did not email a password.",
        ),
        (
            "Work notes",
            "INC-2401: Verified Jordan Lee / jlee / NW-WS-1042. Account locked after failed logons (old password). Unlocked in AD. User confirmed sign-in. No evidence of off-site attack. Closed verified.",
        ),
        ("Outcome", "Verified resolve"),
        ("Skills", "Incident ownership, Active Directory, voice under pressure"),
        (
            "Interview one-liner",
            "I walked a locked-out AP clerk through identity check, unlocked the correct SAM, and documented that it was a bad password — not a security event.",
        ),
    ]
    for label, text in example:
        cells = table.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].paragraphs[0].add_run(text)
    add_p(doc, "", space_after=6)


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("SysDesk — write-up practice template")
    set_run(r, size=22, bold=True, color=NAVY)
    add_p(
        doc,
        "Simulated Northwind MSP campus · help desk, six clients, and the extra floors",
        size=12,
        color=GOLD,
        bold=True,
        space_after=12,
    )

    heading(doc, "How to use this", 1)
    add_p(
        doc,
        "1. Close a ticket in SysDesk (identity, required work, 40+ character notes, Resolve or Escalate).",
    )
    add_p(doc, "2. Copy the blank form (next page). Paste it at the end of this file or into a new doc.")
    add_p(doc, "3. Fill it the same day. If you cannot defend a line out loud, delete the line.")
    add_p(
        doc,
        "4. Bring 4–6 completed pages to interviews as a lab portfolio — not as employment at Northwind or a client.",
    )
    add_p(
        doc,
        "Lead with judgment tickets: Elena rehire, Domain Admin asks, Riverton vs Harbor Maya, Peak vs Oak & Pine Ortiz, ExpressRoute, false-positive lockout, DC patch with no backout.",
        space_after=12,
    )

    heading(doc, "What a good write-up sounds like", 2)
    add_p(
        doc,
        "Name the person, the SAM, the host, and the client. Say what you checked before you changed anything. Say the exact action. Say when you refused. That is the interview.",
    )

    heading(doc, "Ticket cheat sheet (lab move — not a script to memorize)", 1)
    cheat = doc.add_table(rows=1, cols=4)
    cheat.style = "Table Grid"
    for i, h in enumerate(("ID", "Desk", "Case", "Correct move")):
        run = cheat.rows[0].cells[i].paragraphs[0].add_run(h)
        set_run(run, size=9, bold=True, color=NAVY)
    for tid, desk, case, move in TICKETS:
        row = cheat.add_row().cells
        row[0].paragraphs[0].add_run(tid)
        row[1].paragraphs[0].add_run(desk)
        row[2].paragraphs[0].add_run(case)
        row[3].paragraphs[0].add_run(move)
        for cell in row:
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run(run, size=9)

    doc.add_page_break()
    example_form(doc)
    for letter in "ABCDEF":
        doc.add_page_break()
        blank_form(doc, f"Practice write-up {letter} — type in the right column")

    heading(doc, "Resume bullet scratch pad (fill only after verified closes)", 1)
    add_p(
        doc,
        "Practiced ________________________________ in a simulated MSP / enterprise desk, with documented verified resolutions.",
        italic=True,
        color=MUTED,
    )
    add_p(
        doc,
        "Escalated out-of-scope requests (________________________________) instead of making the change.",
        italic=True,
        color=MUTED,
    )
    add_p(
        doc,
        "Kept client directories straight (example: ________________________________).",
        italic=True,
        color=MUTED,
        space_after=16,
    )
    add_p(
        doc,
        "Do not list SysDesk, Northwind, Harbor Dental, or Maple CU as employers. Say simulated lab / portfolio.",
        size=10,
        bold=True,
        color=NAVY,
    )

    doc.save(OUT)
    print(OUT)
    for desktop in (
        Path.home() / "OneDrive" / "Desktop",
        Path.home() / "Desktop",
    ):
        if desktop.is_dir():
            dest = desktop / OUT.name
            dest.write_bytes(OUT.read_bytes())
            print(dest)
            break


if __name__ == "__main__":
    main()
